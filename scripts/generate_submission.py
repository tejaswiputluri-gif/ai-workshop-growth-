from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as SlideRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submissions"
OUT.mkdir(exist_ok=True)

INK = "172C23"
GREEN = "C3EF79"
MINT = "EDF4E5"
PAPER = "F7F8F3"
MUTED = "64746A"
WHITE = "FFFFFF"
LINE = "DCE4DC"


def ppt_rgb(hex_value):
    return SlideRGB.from_string(hex_value)


def add_box(slide, x, y, w, h, color, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        PptInches(x), PptInches(y), PptInches(w), PptInches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = ppt_rgb(color)
    shape.line.fill.background() if not line else None
    if line:
        shape.line.color.rgb = ppt_rgb(line)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def text(slide, value, x, y, w, h, size=12, color=INK, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0):
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = PptInches(margin)
    frame.margin_top = frame.margin_bottom = PptInches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = value
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = PptPt(size)
        run.font.bold = bold
        run.font.color.rgb = ppt_rgb(color)
    return box


def base_slide(prs, number, kicker, title_value, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = ppt_rgb(PAPER)
    add_box(slide, 0, 0, 13.333, 0.12, GREEN)
    text(slide, "N  NXTWAVE  /  GROWTH CHALLENGE", .58, .34, 4.8, .25, 8, MUTED, True, "Aptos")
    text(slide, f"{number:02d}  /  05", 11.6, .34, 1.1, .25, 8, MUTED, True, "Aptos", PP_ALIGN.RIGHT)
    text(slide, kicker.upper(), .62, .83, 6, .25, 8, MUTED, True, "Aptos")
    text(slide, title_value, .62, 1.10, 12, .58, 25, INK, True, "Aptos Display")
    if subtitle:
        text(slide, subtitle, .64, 1.75, 12, .35, 10, MUTED)
    return slide


def make_ppt():
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    s = base_slide(prs, 1, "The student + the job to be done", "From “I studied AI” to “I built this.”",
                   "A credible first project is more motivating than another vague promise about AI.")
    add_box(s, .65, 2.35, 5.75, 3.35, MINT, True)
    text(s, "PRIMARY PERSONA", .95, 2.68, 3, .25, 8, MUTED, True)
    text(s, "Final-year B.Tech\nCSE (AI & ML)", .95, 3.03, 4.9, .92, 26, INK, True, "Aptos Display")
    text(s, "Tier-2/3 college · balancing placements, coursework and project work", .95, 4.05, 4.8, .5, 11, MUTED)
    text(s, "“I know the concepts, but can I build and explain a working project?”", .95, 4.82, 4.85, .62, 13, INK, True)
    text(s, "MOTIVATION", 6.85, 2.53, 2, .22, 8, MUTED, True)
    text(s, "A small guided build they can explain—not a promise of AI mastery or a job.", 6.85, 2.82, 5.4, .57, 14, INK, True)
    text(s, "3 INSIGHTS", 6.85, 3.63, 2, .22, 8, MUTED, True)
    for i, line in enumerate(["Name the project; “free” is not enough.", "Peer context can make a share feel relevant.", "A 60-minute promise needs small scope."]):
        add_box(s, 6.83, 3.94 + i * .48, 5.78, .38, WHITE, True)
        text(s, f"0{i+1}   {line}", 7.02, 4.02 + i * .48, 5.35, .22, 10, INK)
    add_box(s, .65, 6.02, 11.98, .78, INK, True)
    text(s, "VALUE", .93, 6.18, 1, .22, 8, GREEN, True)
    text(s, "Build one small AI chatbot/classifier in 60 minutes—with a guided walkthrough and starter code. (Assumption)", 2.00, 6.15, 10.1, .34, 12, WHITE, True)
    text(s, "ASSUMED EVENT: Saturday 7 PM IST · 60 min · Zoom/Meet + recording · example AI tool/API + starter code", .72, 6.91, 11.9, .18, 7, MUTED, True)

    s = base_slide(prs, 2, "Prioritise trust and trackability", "Four focused channels. One clear ask.",
                   "Permission-based student distribution; projections only, not campaign results.")
    rows = [
        ("01", "College technical clubs", "5,000 reach", "600 / 25%", "150 regs", "₹350", "One approved post with a concrete project and session scope."),
        ("02", "Student ambassadors / peers", "3,500 reach", "525 / 30.1%", "158 regs", "₹450", "Voluntary sharing; unique links and referral codes."),
        ("03", "Permitted WhatsApp / Telegram", "5,000 reach", "500 / 20%", "100 regs", "₹300", "One admin-approved post; no repeated or unsolicited messages."),
        ("04", "Placement communities / email", "7,500 reach", "600 / 15.3%", "92 regs", "₹250", "Relevant communities; email only where permission is appropriate."),
    ]
    headers = ["CHANNEL", "PLANNING REACH", "VISITS · CONV.", "REGISTRATIONS", "COST"]
    widths = [3.15, 1.75, 1.3, 1.55, .9]
    x_positions = [.73, 4.02, 6.04, 7.64, 9.44]
    for x, w, h in zip(x_positions, widths, headers):
        text(s, h, x, 2.28, w, .22, 7, MUTED, True)
    for index, row in enumerate(rows):
        y = 2.58 + index * .83
        add_box(s, .65, y, 11.98, .69, WHITE, True)
        text(s, f"{row[0]}  {row[1]}", .82, y + .14, 3.03, .37, 10, INK, True)
        for x, value, w in zip(x_positions[1:], row[2:6], widths[1:]):
            text(s, value, x, y + .20, w, .22, 10, INK, value == row[4])
        text(s, row[6], 10.6, y + .10, 1.75, .48, 8, MUTED)
    add_box(s, .65, 6.09, 11.98, .7, MINT, True)
    text(s, "MEASURE", .9, 6.31, .9, .2, 8, INK, True)
    text(s, "UTM link per channel + one registration per normalised email. Referral code is a separate inviter field.", 1.85, 6.24, 10.4, .33, 10, INK)

    s = base_slide(prs, 3, "Seven days, ₹2,000 cap", "Planned spend: ₹1,350. Keep ₹650 unspent.",
                   "Simulation only: no campaign spend or outreach has taken place.")
    days = [
        ("D1", "Copy, page,\nlinks"), ("D2", "Ask clubs &\nadmins"), ("D3", "Approved club\nsharing"),
        ("D4", "Voluntary\nreferrals"), ("D5", "Placement /\npermitted email"), ("D6", "Fix friction;\none reminder"),
        ("D7", "Final check;\nreport actuals"),
    ]
    for i, (day, action) in enumerate(days):
        x = .68 + i * 1.78
        add_box(s, x, 2.38, 1.54, 1.04, MINT if i < 5 else WHITE, True)
        text(s, day, x + .13, 2.53, 1.1, .2, 9, INK, True)
        text(s, action, x + .13, 2.82, 1.29, .45, 8, MUTED)
    text(s, "BUDGET LINE", .76, 3.87, 4, .2, 8, MUTED, True)
    text(s, "PLANNED", 5.1, 3.87, 2, .2, 8, MUTED, True)
    budget = [("Club posters / QR material", "₹350"), ("Volunteer coordination / travel", "₹450"),
              ("Community coordination / data", "₹300"), ("Email / material preparation", "₹250"),
              ("Free tools", "₹0"), ("Unallocated buffer", "₹650")]
    for i, (label, amount) in enumerate(budget):
        y = 4.15 + i * .35
        text(s, label, .78, y, 4.2, .2, 9, INK, i == 5)
        text(s, amount, 5.15, y, 1.2, .2, 9, INK, True)
        if i in (3, 5):
            add_box(s, .72, y + .27, 5.9, .008, LINE)
    add_box(s, 7.0, 3.83, 5.55, 2.68, INK, True)
    text(s, "CHECKPOINTS", 7.35, 4.15, 3, .22, 8, GREEN, True)
    text(s, "Day 3 · verify links and visits\nDay 4 · inspect referral attribution\nDay 5 · compare actuals with plan\nDay 6 · improve copy/form before adding channels\nDay 7 · report actuals separately from target", 7.35, 4.53, 4.7, 1.38, 11, WHITE)
    text(s, "Only spend buffer with a clear, permitted use.", 7.35, 6.05, 4.6, .25, 9, GREEN, True)

    s = base_slide(prs, 4, "Transparent planning model", "Target funnel: 21,000 → 2,225 → 500.",
                   "Planning assumptions: reach may overlap; channel-attributed registrations do not.")
    funnel = [("PLANNED REACH", "21,000", "10.6% reach → visit"), ("VISITS", "2,225", "22.5% visit → registration"), ("REGISTRATIONS", "500", "100% of target")]
    start_x = .8
    box_w = 3.72
    for i, (label, num, note) in enumerate(funnel):
        x = start_x + i * 4.15
        add_box(s, x, 2.35, box_w, 1.5, INK if i == 2 else MINT, True)
        text(s, label, x + .25, 2.62, 3.0, .22, 8, GREEN if i == 2 else MUTED, True)
        text(s, num, x + .25, 2.94, 3.0, .42, 28, WHITE if i == 2 else INK, True, "Aptos Display")
        text(s, note, x + .25, 3.47, 3.05, .2, 9, WHITE if i == 2 else MUTED)
        if i < 2:
            text(s, "→", x + 3.75, 2.9, .38, .4, 22, MUTED, True)
    text(s, "REGISTRATION TARGET BY CHANNEL", .78, 4.35, 5, .25, 8, MUTED, True)
    targets = [("Clubs", 150), ("Peers", 158), ("Groups", 100), ("Placement / email", 92)]
    for i, (label, amount) in enumerate(targets):
        y = 4.76 + i * .38
        text(s, label, .83, y, 2.1, .2, 9, INK)
        add_box(s, 2.75, y + .02, 5.2, .14, MINT, True)
        add_box(s, 2.75, y + .02, 5.2 * amount / 158, .14, GREEN, True)
        text(s, str(amount), 8.15, y - .02, .5, .22, 9, INK, True)
    add_box(s, 9.15, 4.45, 3.35, 1.91, WHITE, True)
    text(s, "NO DOUBLE COUNTING", 9.43, 4.72, 2.8, .22, 8, MUTED, True)
    text(s, "One valid email = one registration.\nFirst tracked UTM source = channel.\nReferral code = inviter, tracked separately.\nNo code/source? Keep as direct.", 9.43, 5.09, 2.65, 1.05, 10, INK)
    text(s, "Projection, not guarantee. Do not label these targets as actual sign-ups.", .83, 6.55, 11.3, .2, 9, MUTED, True)

    s = base_slide(prs, 5, "What we can validate now", "A working local prototype—not a live campaign.",
                   "React registration page · FastAPI API · SQLite local storage.")
    items = [
        ("STUDENT FLOW", "Consent, form validation, duplicate email check, unique code, referral URL, WhatsApp share."),
        ("PUBLIC VIEW", "Aggregate count to 500; top referrers by first name/code only."),
        ("ADMIN ONLY", "Token-protected channel/referral/daily metrics, search, CSV and labelled DEMO data."),
    ]
    for i, (label, detail) in enumerate(items):
        y = 2.31 + i * .86
        add_box(s, .7, y, 6.1, .7, WHITE, True)
        text(s, label, .95, y + .13, 1.32, .2, 8, MUTED, True)
        text(s, detail, 2.35, y + .11, 4.12, .43, 9, INK)
    add_box(s, 7.18, 2.3, 5.35, 2.48, MINT, True)
    text(s, "RISKS / NEXT", 7.5, 2.62, 3, .22, 8, MUTED, True)
    text(s, "• Confirm event, project, recording and referral rewards.\n• SQLite is local; do not collect real registrations yet.\n• Before launch: persistent DB, secure admin, durable rate limit, consent/retention review.\n• Below plan? Improve the best converting permitted channel; 500 is a target.", 7.5, 2.98, 4.55, 1.48, 10, INK)
    add_box(s, .7, 5.19, 11.83, .91, INK, True)
    text(s, "PROPOSED REWARDS · ASSUMPTION", .98, 5.38, 2.4, .2, 8, GREEN, True)
    text(s, "3 successful referrals → project template pack  ·  10 → recognition certificate", 3.38, 5.36, 8.6, .28, 11, WHITE, True)
    text(s, "Simulation only. Complete learning notes, rejections, reflections and 3-minute video script are documented in learning-notes.md.", .78, 6.45, 11.8, .22, 9, MUTED)

    prs.save(OUT / "growth-challenge.pptx")


def set_cell_shading(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def cell_text(cell, value, bold=False, color=INK, size=7.4):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def doc_heading(doc, value, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4 if level == 1 else 2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(value)
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(15 if level == 1 else 10)
    run.font.color.rgb = RGBColor.from_string(INK)
    return p


def doc_text(doc, value, size=8.2, color=INK, bold_start=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(value)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return p


def make_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(.43)
    section.bottom_margin = Inches(.42)
    section.left_margin = Inches(.55)
    section.right_margin = Inches(.55)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(8.2)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(1)
    r = title.add_run("NxtWave Growth Challenge")
    r.font.name = "Aptos Display"; r.font.size = Pt(22); r.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(4)
    r = subtitle.add_run("ROUND 1 · SIMULATION PLAN · 7 DAYS · TARGET: 500 REGISTRATIONS")
    r.font.name = "Aptos"; r.font.size = Pt(8); r.bold = True
    r.font.color.rgb = RGBColor.from_string(MUTED)
    doc_text(doc, "No students were contacted, no registrations were acquired and no campaign budget was spent. Reach and conversion figures are projections, not results.", 8, "815B1A")

    doc_heading(doc, "1. Student and strategy")
    doc_text(doc, "Persona: final-year B.Tech CSE (AI & ML) student at a tier-2/3 engineering college, balancing placements, coursework and projects. Interested in AI, but not confident building and explaining a complete project. Likely discovery: technical clubs, peer groups and permitted WhatsApp/Telegram communities. Tier/channel habits are assumptions to validate.")
    doc_text(doc, "Need: “I’ve studied AI, but can I build something that works?” Value: one small guided chatbot/classifier build with starter code—not a promise of AI mastery or a job. Message: “Build a simple AI project in 60 minutes, with a walkthrough and starter code. Free workshop; check the tools and session details before registering.”")
    doc_text(doc, "Insights: (1) “Free” alone is not enough; state the project and scope. (2) Peer context can make a share relevant. (3) A one-hour promise needs a deliberately small build.")

    doc_heading(doc, "2. Prioritised channels and transparent funnel")
    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    widths = [2.4, .76, .68, .84, .6]
    for cell, label, width in zip(table.rows[0].cells, ["Channel / action", "Reach", "Visits", "Regs", "Cost"], widths):
        cell.width = Inches(width); set_cell_shading(cell, INK); cell_text(cell, label, True, WHITE, 7)
    channels = [
        ("Technical clubs · approved club post", "5,000", "600 / 25%", "150", "₹350"),
        ("Student ambassadors · voluntary referrals", "3,500", "525 / 30.1%", "158", "₹450"),
        ("Permitted WhatsApp/Telegram · one approved post", "5,000", "500 / 20%", "100", "₹300"),
        ("Placement communities / permission-based email", "7,500", "600 / 15.3%", "92", "₹250"),
        ("TOTAL projected", "21,000", "2,225 / 22.5%", "500", "₹1,350"),
    ]
    for index, values in enumerate(channels):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            if index == len(channels) - 1: set_cell_shading(cell, MINT)
            cell_text(cell, value, index == len(channels) - 1, INK, 6.8)
    doc_text(doc, "Funnel: 21,000 projected reach → 2,225 visits (10.6%) → 500 registrations (22.5% of visits). Registrations add exactly to 500. Reach may overlap and is not unique verified reach. One valid, trimmed/case-normalised email counts once; first tracked UTM source is the channel; referral code records inviter separately; otherwise source is direct.", 7.6)
    doc_text(doc, "Measurement: unique UTM link by club, ambassador, approved group and placement/email community. Proposed peer rewards (assumptions): 3 successful referrals → template pack; 10 → recognition certificate.", 7.6)

    doc.add_page_break()

    doc_heading(doc, "3. Seven-day runbook")
    daily = [
        ("1", "Finalise message, assumptions, page and source links; test links/form."),
        ("2", "Request club/admin permission; invite voluntary student ambassadors."),
        ("3", "Start approved technical-club and community sharing; check visits/links."),
        ("4", "Start peer referrals; inspect code attribution and reward clarity."),
        ("5", "Share in placement communities; email only where permission is appropriate."),
        ("6", "Improve unclear copy/form; send one reminder only where permitted."),
        ("7", "Final factual reminder where permitted; report actuals separately."),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for day, task in daily:
        cells = table.add_row().cells
        cells[0].width = Inches(.42); cells[1].width = Inches(6.6)
        set_cell_shading(cells[0], GREEN); cell_text(cells[0], f"D{day}", True, INK, 7)
        cell_text(cells[1], task, False, INK, 7.2)

    doc_heading(doc, "4. Budget and response plan")
    budget = doc.add_table(rows=0, cols=2)
    budget.autofit = False
    for item, amount in [
        ("Club posters / QR material", "₹350"), ("Volunteer coordination / travel", "₹450"),
        ("Community coordination / data", "₹300"), ("Email / material preparation", "₹250"),
        ("Free tools", "₹0"), ("Planned spend / unallocated buffer / cap", "₹1,350 / ₹650 / ₹2,000"),
    ]:
        cells = budget.add_row().cells
        cells[0].width = Inches(5.5); cells[1].width = Inches(1.5)
        cell_text(cells[0], item, "cap" in item.lower(), INK, 7.2)
        cell_text(cells[1], amount, True, INK, 7.2)
    doc_text(doc, "Keep the ₹650 buffer unspent unless there is a clear, permitted need. No money has been spent in this simulation.", 7.5, MUTED)

    doc_heading(doc, "5. Working asset, risks and next steps")
    doc_text(doc, "React + FastAPI + local SQLite prototype: consent-based registration, API validation, case-insensitive duplicate-email blocking, referral code/link attribution, WhatsApp share, public aggregate counts/first-name-or-code leaderboard, and admin-token dashboard with channel/referral/daily metrics, search, CSV and clearly labelled demo data.")
    doc_text(doc, "Not deployed; SQLite and in-memory rate limiting are not durable production infrastructure. Do not collect real student data until persistent storage, admin access, backups, consent/retention and organiser approvals are reviewed. Demo records do not count as real; the 500 goal is not guaranteed.")
    doc_text(doc, "Risks: if a group declines, do not post; shift only to approved channels. If visits are low, check link and copy. If registration conversion is low, clarify the exact build and form. If behind target, reallocate time to the best measured permitted channel—not spam or unsupported claims.")
    doc_text(doc, "Assumptions to confirm: Saturday at 7 PM IST, 60 minutes; Zoom/Meet and a recording link; example chatbot/classifier using an AI tool/API and starter code; proposed referral rewards. Full detailed plan, reflections and 3-minute video script are in the companion files.", 7.4, MUTED)

    doc.save(OUT / "growth-plan-two-pages.docx")


def make_pdf():
    path = OUT / "growth-plan-two-pages.pdf"
    doc = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=13 * mm, bottomMargin=14 * mm,
        title="NxtWave Growth Challenge — two-page plan",
        author="Student submission",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="PlanTitle", parent=styles["Title"], fontName="Helvetica-Bold",
                              fontSize=20, leading=23, textColor=colors.HexColor(f"#{INK}"),
                              alignment=TA_LEFT, spaceAfter=2))
    styles.add(ParagraphStyle(name="PlanSub", parent=styles["Normal"], fontName="Helvetica-Bold",
                              fontSize=7.6, leading=9, textColor=colors.HexColor(f"#{MUTED}"),
                              spaceAfter=4))
    styles.add(ParagraphStyle(name="PlanBody", parent=styles["BodyText"], fontName="Helvetica",
                              fontSize=7.6, leading=9.6, textColor=colors.HexColor(f"#{INK}"),
                              spaceAfter=4))
    styles.add(ParagraphStyle(name="PlanSmall", parent=styles["BodyText"], fontName="Helvetica",
                              fontSize=6.9, leading=8.3, textColor=colors.HexColor(f"#{INK}"),
                              spaceAfter=0))
    styles.add(ParagraphStyle(name="PlanHeader", parent=styles["BodyText"], fontName="Helvetica-Bold",
                              fontSize=6.9, leading=8.3, textColor=colors.white, spaceAfter=0))
    styles.add(ParagraphStyle(name="PlanSection", parent=styles["Heading2"], fontName="Helvetica-Bold",
                              fontSize=11, leading=13, textColor=colors.HexColor(f"#{INK}"),
                              spaceBefore=5, spaceAfter=3))
    story = [
        Paragraph("NxtWave Growth Challenge", styles["PlanTitle"]),
        Paragraph("ROUND 1 · SIMULATION PLAN · 7 DAYS · TARGET: 500 REGISTRATIONS", styles["PlanSub"]),
        Paragraph("<b>Simulation only.</b> No students were contacted, registrations acquired or budget spent. Reach and conversion figures are projections, not results.", styles["PlanBody"]),
        Paragraph("1. Student and strategy", styles["PlanSection"]),
        Paragraph("<b>Persona:</b> final-year B.Tech CSE (AI &amp; ML) student at a tier-2/3 engineering college, balancing placements, coursework and projects. Interested in AI, but not confident building and explaining a complete project. Likely discovery: technical clubs, peers and permitted WhatsApp/Telegram groups. College tier and channel habits are assumptions to validate.", styles["PlanBody"]),
        Paragraph("<b>Need:</b> “I’ve studied AI, but can I build something that works?” <b>Value:</b> one small guided chatbot/classifier build with starter code—not a promise of AI mastery or a job. <b>Message:</b> “Build a simple AI project in 60 minutes, with a walkthrough and starter code. Free workshop; check the tools and session details before registering.”", styles["PlanBody"]),
        Paragraph("<b>Insights:</b> (1) “Free” alone is not enough; state the project and scope. (2) Peer context can make a share relevant. (3) A one-hour promise needs a deliberately small build.", styles["PlanBody"]),
        Paragraph("2. Prioritised channels and transparent funnel", styles["PlanSection"]),
    ]
    channel_rows = [
        ["Channel / action", "Reach", "Visits", "Visit→reg", "Regs", "Hours", "Cost"],
        ["Technical clubs · approved post", "5,000", "600", "25%", "150", "5", "₹350"],
        ["Student ambassadors · voluntary referrals", "3,500", "525", "30.1%", "158", "7", "₹450"],
        ["Permitted WhatsApp/Telegram", "5,000", "500", "20%", "100", "4", "₹300"],
        ["Placement communities / appropriate email", "7,500", "600", "15.3%", "92", "5", "₹250"],
        ["TOTAL projected", "21,000", "2,225", "22.5%", "500", "21", "₹1,350"],
    ]
    channel_table = Table(
        [[Paragraph(str(cell), styles["PlanHeader"] if row_index == 0 else styles["PlanSmall"])
          for cell in row] for row_index, row in enumerate(channel_rows)],
        colWidths=[172, 49, 48, 56, 38, 35, 42],
        repeatRows=1,
        hAlign="LEFT",
    )
    channel_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{INK}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor(f"#{MINT}")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor(f"#{LINE}")),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story += [
        channel_table, Spacer(1, 4),
        Paragraph("<b>Funnel:</b> 21,000 projected reach → 2,225 visits (10.6%) → 500 registrations (22.5% of visits). Registrations add exactly to 500. Reach may overlap and is not unique verified reach. One valid, trimmed/case-normalised email counts once; first tracked UTM source is the channel; referral code records inviter separately; otherwise source is direct.", styles["PlanBody"]),
        Paragraph("<b>Measurement:</b> unique UTM link per club, ambassador, approved group and placement/email community. Proposed peer rewards (assumptions): 3 successful referrals → template pack; 10 → recognition certificate.", styles["PlanBody"]),
        PageBreak(),
        Paragraph("Seven-day execution plan", styles["PlanTitle"]),
        Paragraph("3. Runbook", styles["PlanSection"]),
    ]
    daily_rows = [["Day", "Action", "Check"]] + [
        ["1", "Finalise message, assumptions, page and source links.", "Test all links and form."],
        ["2", "Request club/admin permission; invite voluntary ambassadors.", "Record approvals; no unapproved posts."],
        ["3", "Start approved club and community sharing.", "Check visits and broken links."],
        ["4", "Start peer referrals.", "Check attribution and reward clarity."],
        ["5", "Share in relevant placement communities; email only where permission is appropriate.", "Compare actuals with projections."],
        ["6", "Improve unclear copy/form; send one permitted reminder.", "Shift effort to converting channels."],
        ["7", "Final factual reminder where permitted; prepare attendee details.", "Report actuals separately."],
    ]
    daily_table = Table(
        [[Paragraph(str(cell), styles["PlanHeader"] if row_index == 0 else styles["PlanSmall"])
          for cell in row] for row_index, row in enumerate(daily_rows)],
        colWidths=[31, 250, 159], repeatRows=1,
    )
    daily_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{INK}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 1), (0, -1), colors.HexColor(f"#{MINT}")),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor(f"#{LINE}")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story += [daily_table, Paragraph("4. Budget and response plan", styles["PlanSection"])]
    budget_rows = [
        ["Line item", "Planned"],
        ["Club posters / QR material", "₹350"],
        ["Volunteer coordination / travel", "₹450"],
        ["Community coordination / data", "₹300"],
        ["Email / material preparation", "₹250"],
        ["Free tools", "₹0"],
        ["Planned spend / unallocated buffer / cap", "₹1,350 / ₹650 / ₹2,000"],
    ]
    budget_table = Table(
        [[Paragraph(str(cell), styles["PlanHeader"] if row_index == 0 else styles["PlanSmall"])
          for cell in row] for row_index, row in enumerate(budget_rows)],
        colWidths=[340, 100],
    )
    budget_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{INK}")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor(f"#{MINT}")),
        ("GRID", (0, 0), (-1, -1), .35, colors.HexColor(f"#{LINE}")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    story += [
        budget_table, Spacer(1, 3),
        Paragraph("Keep ₹650 unspent unless there is a clear, permitted need. No money has been spent in this simulation.", styles["PlanSmall"]),
        Paragraph("5. Working asset, risks and next steps", styles["PlanSection"]),
        Paragraph("<b>Asset:</b> React + FastAPI + local SQLite prototype: consent-based registration, API validation, case-insensitive duplicate-email blocking, referral code/link attribution, WhatsApp share, public aggregate counts/first-name-or-code leaderboard, and admin-token dashboard with channel/referral/daily metrics, search, CSV and labelled demo data.", styles["PlanBody"]),
        Paragraph("<b>Limits:</b> not deployed; SQLite and in-memory rate limiting are not durable production infrastructure. Do not collect real student data until persistent storage, admin access, backups, consent/retention and organiser approvals are reviewed. Demo records do not count as real; the 500 goal is not guaranteed.", styles["PlanBody"]),
        Paragraph("<b>Response:</b> if a group declines, do not post; shift only to approved channels. If visits are low, check link and copy. If registration conversion is low, clarify the exact build and form. If behind target, reallocate time to the best measured permitted channel—not spam or unsupported claims.", styles["PlanBody"]),
        Paragraph("<b>Assumptions to confirm:</b> Saturday at 7 PM IST, 60 minutes; Zoom/Meet and recording link; example chatbot/classifier using an AI tool/API and starter code; proposed rewards (3 referrals → template pack; 10 → recognition certificate).", styles["PlanBody"]),
    ]

    def page_footer(canvas, _document):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor(f"#{LINE}"))
        canvas.line(15 * mm, 10 * mm, A4[0] - 15 * mm, 10 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawString(15 * mm, 6 * mm, "NXTWAVE GROWTH CHALLENGE · SIMULATION")
        canvas.drawRightString(A4[0] - 15 * mm, 6 * mm, f"PAGE {canvas.getPageNumber()} OF 2")
        canvas.restoreState()

    doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    pages = len(PdfReader(str(path)).pages)
    if pages != 2:
        raise ValueError(f"Expected a two-page plan PDF, generated {pages} pages.")


if __name__ == "__main__":
    make_ppt()
    make_docx()
    make_pdf()
    print(f"Generated {OUT / 'growth-challenge.pptx'}")
    print(f"Generated {OUT / 'growth-plan-two-pages.docx'}")
    print(f"Generated and verified {OUT / 'growth-plan-two-pages.pdf'} (2 pages)")
